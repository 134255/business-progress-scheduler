"""Read a local cached workbook and export an ignored, local-only import sidecar.

Usage: python tools/build-product-option-linkage.py --input SOURCE.xlsx \
    --output qa/catalog.json [--expect-count N] [--expect-digest SHA256]

The import uses logical field keys. Its consumer must bind them to real field
IDs and store optionLinkage on the first member only. This tool never contacts
the application, recalculates formulas, or changes the source workbook.
"""

import argparse
import hashlib
import json
import math
from pathlib import Path
import shutil
import subprocess
import sys
import warnings

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
FIELD_KEYS = ("category", "brand", "model", "attribute1", "attribute2",
              "attribute3", "attribute4", "attribute5")
FIELD_NAMES = ("分类", "品牌", "型号", "属性1", "属性2", "属性3", "属性4", "属性5")
COLUMNS = (11, 3, 1, 4, 5, 6, 7, 8)  # Zero-based: L, D, B, E through I.
ALLOWED_CATEGORIES = frozenset(("灯具", "人体工学椅", "升降桌", "儿童系列"))
EXCLUDED_SOURCE_ROWS = frozenset((2185, 2859, 2860))
MAX_ROWS = 5000
MAX_LINKAGE_BYTES = 256 * 1024
MAX_LINKED_FIELDS_BYTES = 512 * 1024


class ImportValidationError(ValueError):
    """A fixed safe error code, never a source cell value or parser traceback."""


def compact_json(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def normalize_cell(cell):
    if cell.data_type == "e":
        raise ImportValidationError("cell_error")
    if cell.data_type == "f":
        raise ImportValidationError("missing_formula_cache")
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, str):
        # Trim only edges; embedded commas, spaces and newlines are one option.
        return value.strip()
    if type(value) is int or (type(value) is float and math.isfinite(value)):
        return str(value)
    raise ImportValidationError("unsupported_cell_type")


def build_import(workbook, formula_workbook=None):
    """Build from a cached/in-memory workbook, retaining only observed tuples.

    Inconsistent applicability invalidates the entire category/brand/model
    group, not merely the later row. No majority vote or first-row inference.
    Source digest uses trimmed text and empty strings for absent attributes;
    only the indexed output represents those absent attributes with null.
    """
    if "Sheet1" not in workbook.sheetnames:
        raise ImportValidationError("MISSING_SHEET")
    sheet = workbook["Sheet1"]
    formula_rows = None
    if formula_workbook is not None:
        if "Sheet1" not in formula_workbook.sheetnames:
            raise ImportValidationError("MISSING_SHEET")
        formula_rows = iter(formula_workbook["Sheet1"].iter_rows(min_row=2, max_col=12))

    exclusions = []
    candidates = []
    applicability = {}
    source_count = 0
    outside_category_count = 0
    explicit_rows = []
    for number, physical in enumerate(sheet.iter_rows(min_row=2, max_col=12), 2):
        source_count += 1
        formulas = next(formula_rows) if formula_rows is not None else None
        if number in EXCLUDED_SOURCE_ROWS:
            explicit_rows.append(number)
            exclusions.append({"row": number, "reason": "explicit_source_exclusion"})
            continue
        try:
            category = normalize_cell(physical[COLUMNS[0]])
        except ImportValidationError:
            category = ""
        # Count everything outside the allowlist, including an unavailable
        # category cache. Its precise rejection reason remains in exclusions.
        if category not in ALLOWED_CATEGORIES:
            outside_category_count += 1
        try:
            selected = [physical[index] for index in COLUMNS]
            if formulas is not None and any(
                    formulas[index].data_type == "f" and physical[index].value is None
                    for index in COLUMNS):
                raise ImportValidationError("missing_formula_cache")
            record = tuple(normalize_cell(cell) for cell in selected)
        except ImportValidationError as error:
            exclusions.append({"row": number, "reason": str(error)})
            continue
        if not any(record):
            reason = "blank_row"
        elif not record[0]:
            reason = "missing_category"
        elif record[0] not in ALLOWED_CATEGORIES:
            reason = "outside_category"
        elif not record[1]:
            reason = "missing_brand"
        elif not record[2]:
            reason = "missing_model"
        else:
            reason = None
        if reason:
            exclusions.append({"row": number, "reason": reason})
            continue
        group = record[:3]
        applicability.setdefault(group, set()).add(tuple(bool(value) for value in record[3:]))
        candidates.append((number, record))

    tuples = []
    seen = set()
    duplicates = []
    for number, record in candidates:
        if len(applicability[record[:3]]) != 1:
            exclusions.append({"row": number, "reason": "inconsistent_applicability"})
        elif record in seen:
            duplicates.append(number)
        else:
            seen.add(record)
            tuples.append(record)
    if not tuples:
        raise ImportValidationError("NO_VALID_TUPLES")
    if len(tuples) > MAX_ROWS:
        raise ImportValidationError("ROW_LIMIT")

    dictionaries = [[] for _ in FIELD_KEYS]
    indices = [{} for _ in FIELD_KEYS]
    rows = []
    for record in tuples:
        indexed = []
        for column, value in enumerate(record):
            if value == "":
                indexed.append(None)
            else:
                if value not in indices[column]:
                    indices[column][value] = len(dictionaries[column])
                    dictionaries[column].append(value)
                indexed.append(indices[column][value])
        rows.append(indexed)
    for options in dictionaries:
        if not options:
            options.append("不适用")

    fields = [{"fieldKey": key, "name": name, "type": "single_select", "required": True,
               "constraints": {"options": options}}
              for key, name, options in zip(FIELD_KEYS, FIELD_NAMES, dictionaries)]
    linkage = {"schemaVersion": 1, "fieldKeys": list(FIELD_KEYS), "rows": rows}
    linkage_bytes = len(compact_json({"optionLinkage": linkage,
                                     "dictionaries": dictionaries}).encode("utf-8"))
    if linkage_bytes > MAX_LINKAGE_BYTES:
        raise ImportValidationError("LINKAGE_SIZE_LIMIT")
    # Bound the linked-field fragment. The consumer must also bound its complete
    # normalized node (including unrelated fields and generated IDs) at 512 KiB.
    linked_fields = [{**fields[0], "optionLinkage": linkage}, *fields[1:]]
    linked_fields_bytes = len(compact_json({"fields": linked_fields}).encode("utf-8"))
    if linked_fields_bytes > MAX_LINKED_FIELDS_BYTES:
        raise ImportValidationError("LINKED_FIELDS_SIZE_LIMIT")

    exclusions.sort(key=lambda item: item["row"])
    summary = {
        "sourceRowCount": source_count,
        "tupleCount": len(tuples),
        "modelCount": len({record[:3] for record in tuples}),
        "optionCounts": [len(options) for options in dictionaries],
        "sourceTupleDigest": hashlib.sha256(compact_json(tuples).encode("utf-8")).hexdigest(),
        "excludedRowCount": len(exclusions),
        "excludedRows": [item["row"] for item in exclusions],
        "explicitlyExcludedRows": explicit_rows,
        "outsideCategoryCount": outside_category_count,
        "exclusions": exclusions,
        "duplicateRowCount": len(duplicates),
        "duplicateRows": duplicates,
        "linkageBytes": linkage_bytes,
        "linkedFieldsBytes": linked_fields_bytes,
    }
    return {"schemaVersion": 1, "fields": fields, "optionLinkage": linkage, "summary": summary}


def import_workbook(path):
    """Read saved values and formula metadata without calculating or saving."""
    cached = None
    formulas = None
    try:
        # Own the file handles even if the workbook parser fails during load.
        # Parser warning text can contain source values; retain only its count.
        with Path(path).open("rb") as source, Path(path).open("rb") as formula_source:
            with warnings.catch_warnings(record=True) as parser_warnings:
                warnings.simplefilter("always")
                try:
                    cached = load_workbook(source, read_only=True, data_only=True, keep_links=False)
                    formulas = load_workbook(formula_source, read_only=True,
                                             data_only=False, keep_links=False)
                    payload = build_import(cached, formulas)
                    payload["summary"]["workbookWarningCount"] = len(parser_warnings)
                    return payload
                finally:
                    if cached is not None:
                        cached.close()
                    if formulas is not None:
                        formulas.close()
    except ImportValidationError:
        raise
    except Exception:
        # This is the untrusted parser boundary. Do not expose any exception's
        # text (including XML/type errors), which may contain workbook values.
        raise ImportValidationError("WORKBOOK_READ_FAILED") from None


def validate_output_path(path):
    """Resolve links/traversal and require an untracked, Git-ignored qa file."""
    try:
        destination = Path(path).resolve()
        qa = ROOT / "qa"
        relative = destination.relative_to(qa)
        if not relative.parts or destination.suffix.lower() != ".json":
            raise ImportValidationError("OUTPUT_REQUIRES_QA_JSON")
        if any(":" in part for part in relative.parts):
            raise ImportValidationError("OUTPUT_REQUIRES_QA_JSON")
        if destination.exists():
            raise ImportValidationError("OUTPUT_EXISTS")
        # Windows PATH may expose only a shell-wrapper git.cmd. Use a native
        # executable so untrusted path arguments never pass through cmd.exe.
        git = shutil.which("git.exe" if sys.platform == "win32" else "git")
        if not git:
            bundled_git = Path(sys.executable).resolve().parents[1] / "native/git/cmd/git.exe"
            if not bundled_git.is_file():
                raise ImportValidationError("GIT_UNAVAILABLE")
            git = str(bundled_git)
        ignored = subprocess.run(
            [git, "check-ignore", "--quiet", "--", destination.relative_to(ROOT).as_posix()],
            cwd=ROOT, capture_output=True, check=False)
        if ignored.returncode != 0:
            # check-ignore deliberately returns nonzero for a tracked path.
            raise ImportValidationError("OUTPUT_NOT_IGNORED")
        return destination
    except ImportValidationError:
        raise
    except (OSError, ValueError, RuntimeError):
        raise ImportValidationError("OUTPUT_REQUIRES_IGNORED_QA_PATH") from None


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="read-only source workbook")
    parser.add_argument("--output", required=True, type=Path, help="new JSON file inside ignored qa/")
    parser.add_argument("--expect-count", type=int)
    parser.add_argument("--expect-digest", help="SHA-256 of ordered normalized source tuples")
    args = parser.parse_args(argv)
    try:
        destination = validate_output_path(args.output)
        if args.input.resolve() == destination:
            raise ImportValidationError("OUTPUT_IS_SOURCE")
        payload = import_workbook(args.input)
        summary = payload["summary"]
        if args.expect_count is not None and summary["tupleCount"] != args.expect_count:
            raise ImportValidationError("EXPECTED_COUNT_MISMATCH")
        if args.expect_digest is not None and summary["sourceTupleDigest"] != args.expect_digest:
            raise ImportValidationError("EXPECTED_DIGEST_MISMATCH")
        serialized = compact_json(payload)
        destination.parent.mkdir(parents=True, exist_ok=True)
        # Recheck after directory creation; exclusive creation never overwrites a
        # previous artifact, symlink, hard link or the source workbook.
        destination = validate_output_path(args.output)
        with destination.open("x", encoding="utf-8", newline="\n") as output:
            output.write(serialized + "\n")
        print(compact_json(summary))
        return 0
    except ImportValidationError as error:
        print(f"Import failed: {error}", file=sys.stderr)
    except OSError:
        print("Import failed: LOCAL_IO_FAILED", file=sys.stderr)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
