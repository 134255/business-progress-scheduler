"""Build the staff guide from its reviewed Markdown source, without cloud writes."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/user-guides/45minup售后首次使用说明.md'
OUTPUT = ROOT / 'outputs/user-guides/45minup售后小程序首次使用说明.docx'


def font(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for script in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + script), 'Microsoft YaHei')


def table(doc, lines):
    rows = [[x.strip() for x in line.strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', c) for c in r)]
    count = len(rows[0])
    widths = [1.4, 5.6] if count == 2 else [1.8, 1.0, 4.2]
    tbl = doc.add_table(rows=0, cols=count)
    tbl.autofit = False
    for c, width in zip(tbl.columns, widths):
        c.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + k), v)
        borders.append(e)
    tbl._tbl.tblPr.append(borders)
    for i, row in enumerate(rows):
        cells = tbl.add_row().cells
        pr = tbl.rows[-1]._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            pr.append(OxmlElement('w:tblHeader'))
        for c, text, width in zip(cells, row, widths):
            c.width = Inches(width)
            c.vertical_alignment = 1
            margin = OxmlElement('w:tcMar')
            for edge, value in [('top','65'),('bottom','65'),('left','90'),('right','90')]:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), value)
                e.set(qn('w:type'), 'dxa')
                margin.append(e)
            c._tc.get_or_add_tcPr().append(margin)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            run = p.add_run(text)
            run.font.size = Pt(10.5)
            run.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.63), Inches(.62)
    sec.left_margin, sec.right_margin = Inches(.75), Inches(.75)
    sec.header_distance, sec.footer_distance = Inches(.25), Inches(.25)
    font(doc.styles['Normal'], 11)
    normal = doc.styles['Normal'].paragraph_format
    normal.space_after, normal.line_spacing = Pt(5), 1.12
    font(doc.styles['Title'], 22, True)
    font(doc.styles['Heading 1'], 16, True)
    font(doc.styles['Heading 2'], 12, True)
    for name in ('Title', 'Heading 1', 'Heading 2'):
        fmt = doc.styles[name].paragraph_format
        fmt.space_before = Pt(8 if name == 'Heading 2' else 0)
        fmt.space_after = Pt(7)
        fmt.keep_with_next = True
    for name in ('List Bullet', 'List Number'):
        font(doc.styles[name], 11)
        doc.styles[name].paragraph_format.space_after = Pt(4)
    header = sec.header.paragraphs[0]
    run = header.add_run('45minup售后 · 员工首次使用说明')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string('666666')
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('模板第13版 · 第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')
    for run in footer.runs:
        run.font.size = Pt(9)
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == '<!-- pagebreak -->':
            doc.add_page_break()
        elif line.startswith('|'):
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip())
                i += 1
            table(doc, block)
            continue
        elif line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            doc.add_paragraph(line[3:], 'Heading 1')
        elif line.startswith('### '):
            doc.add_paragraph(line[4:], 'Heading 2')
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], 'List Bullet')
        else:
            p = doc.add_paragraph(line)
            if re.match(r'^\d+\. ', line):
                p.paragraph_format.left_indent = Inches(.18)
                p.paragraph_format.first_line_indent = Inches(-.18)
        i += 1
    doc.core_properties.title = '45minup售后小程序首次使用说明'
    doc.core_properties.subject = '启用模板第13版员工操作手册'
    doc.core_properties.author = '45minup'
    doc.core_properties.keywords = '首次使用,售后,商品联动,处理,审核'
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    check = Document(OUTPUT)
    assert len(check.tables) == 5
    assert sum('type="page"' in p._p.xml for p in check.paragraphs) == 5
    all_text = '\n'.join(p.text for p in check.paragraphs) + '\n' + '\n'.join(c.text for t in check.tables for r in t.rows for c in r.cells)
    for line in lines:
        text = line.strip()
        if not text or text.startswith('<!--'):
            continue
        if text.startswith('|'):
            parts = [part.strip() for part in text.strip('|').split('|')]
            if all(re.fullmatch(r':?-+:?', part) for part in parts):
                continue
        else:
            parts = [re.sub(r'^(?:#{1,3} |\- )', '', text)]
        for part in parts:
            assert part in all_text, f'Missing source text: {part[:40]}'
    for title in ('售后信息收集', '售后解决方案提供', '退货方案执行', '退货单号提供',
                  '换货方案执行', '换货退回', '更换配件方案执行', '配件退还', '赠品补偿方案执行', '其他方案执行'):
        assert title in all_text, title
    assert '尚未接通个人微信订阅消息' in all_text
    assert '保存处理进度' in all_text and '发送 CSV' in all_text
    assert not re.search(r'cloud1-|user_[a-f0-9]{20}|template_[a-f0-9]{20}', all_text)
    print(f'Created: {OUTPUT}')
    print(f'Structural check: {len(check.paragraphs)} paragraphs, {len(check.tables)} tables, 5 explicit page breaks; source-text parity and all 10 node names verified')


if __name__ == '__main__':
    main()
